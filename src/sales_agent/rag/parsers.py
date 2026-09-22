from __future__ import annotations

import asyncio
import csv
import io
import json
from collections.abc import Callable
from pathlib import Path


class DocumentParseError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("TEXT_ENCODING_UNSUPPORTED")


def _parse_text(data: bytes) -> str:
    return _decode(data)


def _parse_json(data: bytes) -> str:
    try:
        value = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise DocumentParseError("JSON_INVALID") from exc
    return json.dumps(value, ensure_ascii=False, indent=2)


def _parse_csv(data: bytes) -> str:
    reader = csv.reader(io.StringIO(_decode(data)))
    return "\n".join("\t".join(cell.strip() for cell in row) for row in reader)


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentParseError("PDF_PARSER_NOT_INSTALLED") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentParseError("PDF_PARSE_FAILED") from exc


def _parse_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParseError("DOCX_PARSER_NOT_INSTALLED") from exc
    try:
        document = Document(io.BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        tables = [
            "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
            for table in document.tables
        ]
        return "\n\n".join([*paragraphs, *tables])
    except Exception as exc:
        raise DocumentParseError("DOCX_PARSE_FAILED") from exc


def _parse_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("XLSX_PARSER_NOT_INSTALLED") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = []
        for worksheet in workbook.worksheets:
            rows = [
                "\t".join("" if cell is None else str(cell) for cell in row)
                for row in worksheet.iter_rows(values_only=True)
            ]
            sheets.append(f"# {worksheet.title}\n" + "\n".join(rows))
        workbook.close()
        return "\n\n".join(sheets)
    except Exception as exc:
        raise DocumentParseError("XLSX_PARSE_FAILED") from exc


_PARSERS: dict[str, Callable[[bytes], str]] = {
    ".txt": _parse_text,
    ".md": _parse_text,
    ".markdown": _parse_text,
    ".json": _parse_json,
    ".csv": _parse_csv,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
}


async def parse_document(filename: str, data: bytes, *, max_chars: int = 2_000_000) -> str:
    suffix = Path(filename).suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise DocumentParseError("DOCUMENT_TYPE_UNSUPPORTED")
    text = (await asyncio.to_thread(parser, data)).strip()
    if not text:
        raise DocumentParseError("DOCUMENT_EMPTY_AFTER_PARSE")
    if len(text) > max_chars:
        raise DocumentParseError("EXTRACTED_TEXT_TOO_LARGE")
    return text
